from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.analyzer.metrics import AnalysisResult

# ── Palette ──────────────────────────────────────────────────────────────────
HEADER_FILL = PatternFill(start_color="2E4057", end_color="2E4057", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
ACCENT_FILL = PatternFill(start_color="048A81", end_color="048A81", fill_type="solid")
ACCENT_FONT = Font(color="FFFFFF", bold=True, size=11)


def write_dataframe_to_sheet(ws, df: pd.DataFrame, start_row: int = 1) -> None:
    """Write a DataFrame into a worksheet starting at the given row.

    The first row written contains the column headers.
    Values are written as plain Python types so openpyxl can serialize them.
    """
    for col_idx, col_name in enumerate(df.columns, start=1):
        ws.cell(row=start_row, column=col_idx, value=col_name)

    for row_idx, row in enumerate(df.itertuples(index=False), start=start_row + 1):
        for col_idx, value in enumerate(row, start=1):
            if hasattr(value, "item"):
                value = value.item()
            ws.cell(row=row_idx, column=col_idx, value=value)


def auto_fit_columns(ws, min_width: int = 12, max_width: int = 40) -> None:
    """Adjust each column width to fit its longest value."""
    for col_cells in ws.columns:
        max_len = max(
            (len(str(cell.value)) for cell in col_cells if cell.value is not None),
            default=min_width,
        )
        col_letter = get_column_letter(col_cells[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, min_width), max_width)


def apply_header_style(ws, header_row: int, col_count: int) -> None:
    """Apply bold white text on a dark background to the header row."""
    for col_idx in range(1, col_count + 1):
        cell = ws.cell(row=header_row, column=col_idx)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")


def create_summary_sheet(ws, result: AnalysisResult, df: pd.DataFrame) -> None:
    """Write the KPI summary block into the given worksheet."""
    best_product = result.by_product.iloc[0]["product"]
    best_region = result.by_region.iloc[0]["region"]
    total_orders = int(df.shape[0])
    total_units = int(df["quantity"].sum())
    avg_order_value = round(df["total"].mean(), 2)

    kpis = [
        ("Metric", "Value"),
        ("Total Revenue", f"${result.total_revenue:,.2f}"),
        ("Total Orders", total_orders),
        ("Total Units Sold", total_units),
        ("Avg Order Value", f"${avg_order_value:,.2f}"),
        ("Best Product", best_product),
        ("Best Region", best_region),
    ]

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 22

    title_cell = ws.cell(row=1, column=1, value="Sales Report — Summary")
    title_cell.font = Font(bold=True, size=14, color="FFFFFF")
    title_cell.fill = ACCENT_FILL
    title_cell.alignment = Alignment(horizontal="center")
    ws.merge_cells("A1:B1")

    for row_idx, (metric, value) in enumerate(kpis, start=3):
        label = ws.cell(row=row_idx, column=1, value=metric)
        val = ws.cell(row=row_idx, column=2, value=value)
        if row_idx == 3:
            label.font = HEADER_FONT
            label.fill = HEADER_FILL
            val.font = HEADER_FONT
            val.fill = HEADER_FILL
        label.alignment = Alignment(horizontal="left")
        val.alignment = Alignment(horizontal="right")
